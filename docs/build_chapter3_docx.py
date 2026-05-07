"""Convert chapter3.md and annotation.md to Word documents."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_font(run, size=14, bold=False, italic=False, name='Times New Roman'):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._r.get_or_add_rPr()
    for tag in [qn('w:rFonts')]:
        el = rpr.find(tag)
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn('w:ascii'), name)
        el.set(qn('w:hAnsi'), name)
        el.set(qn('w:cs'), name)


def set_para_fmt(para, indent=True, space_before=0, space_after=0,
                 align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = para.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Cm(1.25) if indent else Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def add_heading(doc, text, level):
    para = doc.add_paragraph()
    set_para_fmt(para, indent=False,
                 space_before=14 if level <= 2 else 8,
                 space_after=6,
                 align=WD_ALIGN_PARAGRAPH.LEFT)
    run = para.add_run(text)
    set_font(run, size=14, bold=True)
    return para


def add_body(doc, text):
    para = doc.add_paragraph()
    set_para_fmt(para)
    parts = re.split(r'`([^`]+)`', text)
    for i, part in enumerate(parts):
        if not part:
            continue
        run = para.add_run(part)
        if i % 2 == 1:
            set_font(run, size=12, name='Courier New')
        else:
            set_font(run, size=14)
    return para


def add_bullet(doc, text):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = Cm(1.25)
    pf.first_line_indent = Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = para.add_run(text)
    set_font(run, size=14)
    return para


def add_numbered(doc, text):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = Cm(1.25)
    pf.first_line_indent = Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = para.add_run(text)
    set_font(run, size=14)
    return para


def add_table_from_md(doc, rows):
    if len(rows) < 2:
        return
    # Remove separator row (---|---|---)
    data_rows = [r for r in rows if not re.match(r'^[\|\s\-:]+$', r)]
    if not data_rows:
        return
    parsed = []
    for row in data_rows:
        cells = [c.strip() for c in row.strip('|').split('|')]
        parsed.append(cells)
    if not parsed:
        return
    col_count = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=col_count)
    table.style = 'Table Grid'
    for i, row_data in enumerate(parsed):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            if j < col_count:
                cell = row.cells[j]
                cell.text = ''
                para = cell.paragraphs[0]
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                run = para.add_run(cell_text)
                set_font(run, size=12, bold=(i == 0))
    # Add spacing after table
    doc.add_paragraph()


def init_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    return doc


def build_doc(md_path: Path, out_path: Path):
    doc = init_doc()
    lines = md_path.read_text(encoding='utf-8').splitlines()

    table_buf = []
    in_table = False

    for line in lines:
        stripped = line.strip()

        # Table handling
        if stripped.startswith('|'):
            in_table = True
            table_buf.append(stripped)
            continue
        else:
            if in_table:
                add_table_from_md(doc, table_buf)
                table_buf = []
                in_table = False

        if not stripped or stripped == '---':
            continue

        if stripped.startswith('# '):
            add_heading(doc, stripped[2:], 1)
        elif stripped.startswith('## '):
            add_heading(doc, stripped[3:], 2)
        elif stripped.startswith('### '):
            add_heading(doc, stripped[4:], 3)
        elif re.match(r'^\d+\.', stripped):
            add_numbered(doc, stripped)
        elif stripped.startswith('– ') or stripped.startswith('- '):
            add_bullet(doc, stripped)
        elif stripped.startswith('**') and stripped.endswith('**'):
            para = doc.add_paragraph()
            set_para_fmt(para, indent=True)
            run = para.add_run(stripped.strip('*'))
            set_font(run, size=14, bold=True)
        else:
            add_body(doc, stripped)

    if in_table and table_buf:
        add_table_from_md(doc, table_buf)

    doc.save(str(out_path))
    print(f'Saved: {out_path}')


if __name__ == '__main__':
    base = Path(__file__).parent
    build_doc(base / 'chapter3.md', base / 'chapter3.docx')
    build_doc(base / 'annotation.md', base / 'annotation.docx')
